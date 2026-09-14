"""
Resume text extraction (spec section 6: "Resume analysis" pipeline starts
with getting plain text out of whatever file the user uploaded).

Deliberately synchronous, CPU-bound functions — they're called from inside
the background task (resume_service.process_resume_upload), never directly
from an `async def` route handler, so blocking here doesn't stall the event
loop for other requests.
"""
from app.core.errors import CorruptFileError


def extract_text(file_path: str, file_type: str) -> str:
    """Returns the resume's plain text. Raises CorruptFileError if the file
    can't actually be parsed as the type its extension claims."""
    if file_type == "pdf":
        return _extract_pdf_text(file_path)
    if file_type == "docx":
        return _extract_docx_text(file_path)
    raise ValueError(f"Unsupported file_type for extraction: {file_type}")


def _extract_pdf_text(file_path: str) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(file_path)
        pages_text = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, OSError, ValueError) as exc:
        raise CorruptFileError() from exc

    return "\n".join(pages_text)


def _extract_docx_text(file_path: str) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(file_path)
    except (PackageNotFoundError, KeyError, OSError) as exc:
        raise CorruptFileError() from exc

    paragraphs = [p.text for p in document.paragraphs]

    # python-docx doesn't walk table cells as part of .paragraphs — a lot of
    # resume templates put the Skills or Contact section in a table, so
    # skipping this would silently drop real content.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)
